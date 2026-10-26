import csv
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLUMNS = 60


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    if len(text) > 2000:
        raise ValidationError("Buňka obsahuje příliš dlouhou hodnotu (maximum 2000 znaků).")
    return text


def parse_upload(file, sheet=None):
    data = file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise ValidationError("Soubor může mít nejvýše 5 MiB.")
    suffix = Path(file.name).suffix.lower()
    try:
        if suffix == ".csv":
            text = data.decode("utf-8-sig")
            try:
                delimiter = csv.Sniffer().sniff(text[:8192], delimiters=";,\t").delimiter
            except csv.Error:
                delimiter = ";"
            iterator = csv.reader(StringIO(text), delimiter=delimiter)
            rows = []
            for row in iterator:
                if len(rows) > MAX_ROWS:
                    raise ValidationError("Import může mít nejvýše 10000 datových řádků.")
                if len(row) > MAX_COLUMNS:
                    raise ValidationError("Import může mít nejvýše 60 sloupců.")
                rows.append([_cell(v) for v in row])
        elif suffix == ".xlsx":
            with ZipFile(BytesIO(data)) as archive:
                items = archive.infolist()
                if (
                    len(items) > 2000
                    or sum(i.file_size for i in items) > 50 * 1024 * 1024
                    or any("vbaProject" in i.filename for i in items)
                ):
                    raise ValidationError(
                        "Sešit má příliš velký rozbalený obsah nebo obsahuje makra."
                    )
            raw = load_workbook(BytesIO(data), read_only=True, data_only=False)
            cached = load_workbook(BytesIO(data), read_only=True, data_only=True)
            try:
                if not sheet and len(raw.sheetnames) != 1:
                    raise ValidationError(
                        "Sešit má více listů. Nahrajte jej znovu a vyberte list: "
                        + ", ".join(raw.sheetnames)
                    )
                sheet = sheet or raw.sheetnames[0]
                if sheet not in raw.sheetnames:
                    raise ValidationError("Vybraný list v sešitu neexistuje.")
                ws = raw[sheet]
                cs = cached[sheet]
                if ws.max_column and ws.max_column > MAX_COLUMNS:
                    raise ValidationError("Import může mít nejvýše 60 sloupců.")
                rows = []
                for index, (raw_row, cached_row) in enumerate(zip(ws.iter_rows(), cs.iter_rows())):
                    if index > MAX_ROWS:
                        raise ValidationError("Import může mít nejvýše 10000 datových řádků.")
                    values = []
                    for original, value in zip(raw_row, cached_row):
                        if original.data_type == "f" and value.value is None:
                            raise ValidationError(
                                f"Řádek {index + 1}: vzorec nemá uložený výsledek. V Excelu přepočítejte a uložte sešit nebo vložte hodnoty."
                            )
                        values.append(_cell(value.value))
                    rows.append(values)
            finally:
                raw.close()
                cached.close()
        else:
            raise ValidationError("Podporovány jsou soubory .xlsx a .csv v UTF-8.")
    except ValidationError:
        raise
    except (
        UnicodeDecodeError,
        csv.Error,
        BadZipFile,
        ValueError,
        KeyError,
        OSError,
        OverflowError,
        ParseError,
        DefusedXmlException,
        InvalidFileException,
    ) as exc:
        raise ValidationError("Soubor nelze přečíst. Ověřte formát, kódování a obsah.") from exc
    if len(rows) < 2:
        raise ValidationError("Soubor musí obsahovat hlavičku a alespoň jeden datový řádek.")
    headers = rows[0]
    if not headers or any(not h for h in headers) or len(headers) != len(set(headers)):
        raise ValidationError("Každý sloupec musí mít vyplněnou jedinečnou hlavičku.")
    if any(len(row) > len(headers) for row in rows[1:]):
        raise ValidationError("Datový řádek má více hodnot než hlavička.")
    return {"headers": headers, "rows": rows[1:], "sheet": sheet if suffix == ".xlsx" else ""}, data
