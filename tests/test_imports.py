from datetime import timedelta
from io import BytesIO

import pytest
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone

pytestmark = pytest.mark.django_db
MAPPING = {
    "date": "Datum",
    "user": "Externista",
    "project": "Projekt",
    "hours": "Hodiny",
    "description": "Popis",
}


def csv_file(rows="2026-09-01;jana;WEB;1,5;Analýza\n", name="history.csv"):
    return SimpleUploadedFile(
        name, ("Datum;Externista;Projekt;Hodiny;Popis\n" + rows).encode("utf-8-sig")
    )


def test_preview_writes_no_reports_and_confirm_is_idempotent(world):
    from imports.services import confirm_batch, prepare_batch
    from reports.models import MonthlyReport

    batch = prepare_batch(actor=world["manager"], file=csv_file(), mapping=MAPPING)
    assert not batch.errors
    assert MonthlyReport.objects.count() == 1
    result = confirm_batch(actor=world["manager"], batch_id=batch.pk)
    assert result["reports"] == 1
    report = MonthlyReport.objects.get(month="2026-09-01")
    assert report.status == "submitted"
    assert report.entries.get().minutes == 90
    assert confirm_batch(actor=world["manager"], batch_id=batch.pk) == result
    assert MonthlyReport.objects.count() == 2


@pytest.mark.parametrize(
    "rows",
    [
        "2026-09-01;jana;WEB;5;Práce\n2026-09-01;jana;DATA;4;Práce\n",
        "2026-09-01;jana;WEB;1;Práce\n2026-09-02;unknown;WEB;1;Práce\n",
        "2026-10-01;jana;WEB;1;Existující měsíc\n",
        "2026-09-01;jana;WEB;-1;Chyba\n",
        "2026-09-01;jana;WEB;0,001;Neúplná minuta\n",
    ],
)
def test_invalid_batch_cannot_partially_import(world, rows):
    from imports.services import confirm_batch, prepare_batch
    from reports.models import MonthlyReport

    batch = prepare_batch(actor=world["manager"], file=csv_file(rows), mapping=MAPPING)
    assert batch.errors
    with pytest.raises(ValidationError):
        confirm_batch(actor=world["manager"], batch_id=batch.pk)
    assert MonthlyReport.objects.count() == 1


def test_confirm_revalidates_changed_terms(world):
    from imports.services import confirm_batch, prepare_batch

    batch = prepare_batch(actor=world["manager"], file=csv_file(), mapping=MAPPING)
    terms = world["assignments"][0].terms.get()
    terms.monthly_limit_minutes = 60
    terms.save()
    with pytest.raises(ValidationError):
        confirm_batch(actor=world["manager"], batch_id=batch.pk)


def test_duplicate_rows_require_explicit_confirmation(world):
    from imports.services import confirm_batch, prepare_batch

    batch = prepare_batch(
        actor=world["manager"], file=csv_file("2026-09-01;jana;WEB;1;Stejné\n" * 2), mapping=MAPPING
    )
    assert batch.duplicates
    with pytest.raises(ValidationError):
        confirm_batch(actor=world["manager"], batch_id=batch.pk)
    assert (
        confirm_batch(actor=world["manager"], batch_id=batch.pk, accept_duplicates=True)["entries"]
        == 2
    )


def test_expired_or_foreign_batch_rejected(world):
    from imports.services import confirm_batch, prepare_batch

    batch = prepare_batch(actor=world["manager"], file=csv_file(), mapping=MAPPING)
    world["other"].is_staff = True
    world["other"].save()
    with pytest.raises(PermissionDenied):
        confirm_batch(actor=world["other"], batch_id=batch.pk)
    batch.expires_at = timezone.now() - timedelta(seconds=1)
    batch.save()
    with pytest.raises(ValidationError):
        confirm_batch(actor=world["manager"], batch_id=batch.pk)


def test_xlsx_dates_and_missing_formula_cache(world):
    from datetime import date

    from openpyxl import Workbook

    from imports.services import prepare_batch

    wb = Workbook()
    ws = wb.active
    ws.append(["Datum", "Externista", "Projekt", "Hodiny", "Popis"])
    ws.append([date(2026, 9, 1), "jana", "WEB", 1.5, "Analýza"])
    out = BytesIO()
    wb.save(out)
    batch = prepare_batch(
        actor=world["manager"],
        file=SimpleUploadedFile("history.xlsx", out.getvalue()),
        mapping=MAPPING,
    )
    assert batch.rows[0]["minutes"] == 90
    ws["D2"] = "=1+1"
    out = BytesIO()
    wb.save(out)
    with pytest.raises(ValidationError):
        prepare_batch(
            actor=world["manager"],
            file=SimpleUploadedFile("formula.xlsx", out.getvalue()),
            mapping=MAPPING,
        )


def test_completed_file_fingerprint_rejected(world):
    from imports.services import confirm_batch, prepare_batch

    batch = prepare_batch(actor=world["manager"], file=csv_file(), mapping=MAPPING)
    confirm_batch(actor=world["manager"], batch_id=batch.pk)
    with pytest.raises(ValidationError):
        prepare_batch(actor=world["manager"], file=csv_file(), mapping=MAPPING)


def test_oversized_and_old_binary_rejected(world):
    from imports.services import prepare_batch

    for file in [
        SimpleUploadedFile("old.xls", b"old"),
        SimpleUploadedFile("large.csv", b"x" * (5 * 1024 * 1024 + 1)),
    ]:
        with pytest.raises(ValidationError):
            prepare_batch(actor=world["manager"], file=file, mapping=MAPPING)


def test_malformed_xlsx_xml_is_reported_as_validation_error(world):
    from io import BytesIO
    from zipfile import ZipFile

    from openpyxl import Workbook

    from imports.services import prepare_batch

    wb = Workbook()
    wb.active.append(["Datum", "Externista", "Projekt", "Hodiny"])
    out = BytesIO()
    wb.save(out)
    broken = BytesIO()
    with ZipFile(BytesIO(out.getvalue())) as source, ZipFile(broken, "w") as target:
        for item in source.infolist():
            target.writestr(
                item,
                b"<broken"
                if item.filename == "[Content_Types].xml"
                else source.read(item.filename),
            )
    with pytest.raises(ValidationError):
        prepare_batch(
            actor=world["manager"],
            file=SimpleUploadedFile("broken.xlsx", broken.getvalue()),
            mapping=MAPPING,
        )


def test_zip_bomb_rejected_before_parsing(world):
    from io import BytesIO
    from zipfile import ZIP_DEFLATED, ZipFile

    from imports.services import prepare_batch

    data = BytesIO()
    with ZipFile(data, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("oversized.xml", b"0" * (50 * 1024 * 1024 + 1))
    with pytest.raises(ValidationError):
        prepare_batch(
            actor=world["manager"],
            file=SimpleUploadedFile("oversized.xlsx", data.getvalue()),
            mapping=MAPPING,
        )


def test_distinct_worksheets_can_be_imported_separately(world):
    from io import BytesIO

    from openpyxl import Workbook

    from imports.services import confirm_batch, prepare_batch
    from reports.models import MonthlyReport

    wb = Workbook()
    september = wb.active
    september.title = "September"
    august = wb.create_sheet("August")
    for sheet, day in [(september, "2026-09-01"), (august, "2026-08-01")]:
        sheet.append(["Datum", "Externista", "Projekt", "Hodiny", "Popis"])
        sheet.append([day, "jana", "WEB", 1, "Práce"])
    data = BytesIO()
    wb.save(data)
    first = prepare_batch(
        actor=world["manager"],
        file=SimpleUploadedFile("history.xlsx", data.getvalue()),
        mapping=MAPPING,
        sheet="September",
    )
    confirm_batch(actor=world["manager"], batch_id=first.pk)
    second = prepare_batch(
        actor=world["manager"],
        file=SimpleUploadedFile("history.xlsx", data.getvalue()),
        mapping=MAPPING,
        sheet="August",
    )
    confirm_batch(actor=world["manager"], batch_id=second.pk)
    assert MonthlyReport.objects.filter(user=world["person"], status="submitted").count() == 2
    with pytest.raises(ValidationError):
        prepare_batch(
            actor=world["manager"],
            file=SimpleUploadedFile("history.xlsx", data.getvalue()),
            mapping=MAPPING,
            sheet="September",
        )
